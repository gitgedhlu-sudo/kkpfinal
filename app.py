from flask import Flask, render_template, request, jsonify, send_file, abort, session, redirect, url_for
from datetime import datetime
from pathlib import Path
import json
import os
import sqlite3
from functools import wraps
import secrets
import re
import shutil
from urllib.parse import quote_plus
from io import BytesIO
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment

BASE_DIR = Path(__file__).resolve().parent
PRODUCTS_FILE = BASE_DIR / "products.json"
# On Railway, attach a Volume at /app/data. For local development, data is stored
# in ./data unless ORDER_DATA_DIR is explicitly set.
DATA_DIR = Path(os.environ.get("ORDER_DATA_DIR", str(BASE_DIR / "data"))).resolve()
ORDERS_DIR = DATA_DIR / "orders"
ALL_ORDERS_DIR = DATA_DIR / "all_orders"
DB_FILE = DATA_DIR / "orders.sqlite3"
DATA_DIR.mkdir(parents=True, exist_ok=True)
ORDERS_DIR.mkdir(parents=True, exist_ok=True)
ALL_ORDERS_DIR.mkdir(parents=True, exist_ok=True)

app = Flask(__name__, template_folder="templates", static_folder="public", static_url_path="")
app.secret_key = os.environ.get("FLASK_SECRET_KEY") or secrets.token_hex(32)
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax")
if os.environ.get("RAILWAY_PUBLIC_DOMAIN") or os.environ.get("RAILWAY_ENVIRONMENT"):
    app.config["SESSION_COOKIE_SECURE"] = True

# Initial credentials requested by the site owner. Override via Railway Variables
# before public launch; passwords are configurable without editing source code.
LOGIN_USERS = {
    os.environ.get("LOGIN_ADMIN_USER", "admin"): {"password": os.environ.get("LOGIN_ADMIN_PASSWORD", "dhavalpatel"), "role": "admin"},
    os.environ.get("LOGIN_USER_1", "navsariyuvakmandalpurv"): {"password": os.environ.get("LOGIN_PASSWORD_1", "Navsari@Kkp26!"), "role": "customer"},
    os.environ.get("LOGIN_USER_2", "suratyuvakmandal"): {"password": os.environ.get("LOGIN_PASSWORD_2", "Surat#Kkp84!"), "role": "customer"},
    os.environ.get("LOGIN_USER_3", "valsadyuvukmandal"): {"password": os.environ.get("LOGIN_PASSWORD_3", "Valsad@Kkp73!"), "role": "customer"},
    os.environ.get("LOGIN_USER_4", "vesmayuvukmandal"): {"password": os.environ.get("LOGIN_PASSWORD_4", "vesma@Kkp26!"), "role": "customer"},
}

def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("username"):
            return redirect(url_for("login", next=request.path))
        if session.get("role") != "admin":
            abort(403)
        return view(*args, **kwargs)
    return wrapped


def style_workbook(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="172554")
        cell.alignment = Alignment(vertical="center")
    # Use the numeric column index rather than cell.column_letter: cells inside
    # merged ranges (e.g. B1:K1 when A1:K1 is merged) are MergedCell objects
    # and do not expose column_letter on all openpyxl versions.
    for col_idx, column in enumerate(ws.iter_cols(), start=1):
        letter = get_column_letter(col_idx)
        max_len = max((len(str(c.value or "")) for c in column), default=10)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 42)


def build_order_excel(order, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Order Receipt"
    ws.append(["KKP FIREWORKS — ORDER RECEIPT"])
    ws.append(["Order number", order["order_id"]])
    ws.append(["Date", order["created_at"].replace("T", " ")])
    ws.append(["Login account", order.get("username", "unknown")])
    ws.append(["Mandal Name", order.get("mandal_name", order["customer_name"])])
    ws.append(["Firm Name", order.get("firm_name", "Not provided")])
    ws.append(["Mobile No", order.get("mobile_no", order.get("phone") or "Not provided")])
    ws.append([])
    ws.append(["Product code", "Product name", "Quantity", "Unit price (INR)", "Amount (INR)"])
    for item in order["items"]:
        ws.append([item["code"], item["name"], int(item["quantity"]), int(item["unit_price"]), int(item["line_total"])])
    ws.append(["", "", "", "TOTAL", int(order["total"])])
    ws.append([])
    ws.append(["Order status", order.get("status", "Order request received")])
    ws.append(["Created by", "Aditya Maganbhai Vasani"])
    ws.merge_cells("A1:E1")
    ws["A1"].font = Font(bold=True, size=15, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="172554")
    ws["A1"].alignment = Alignment(horizontal="center")
    for c in ws[9]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="172554")
    for c in ws[ws.max_row-3]:
        pass
    for row in ws.iter_rows(min_row=10, max_row=ws.max_row):
        for c in row:
            if c.column in (4,5) and isinstance(c.value, (int,float)):
                c.number_format = '₹#,##0'
    for col, width in {"A":18,"B":38,"C":14,"D":20,"E":20}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A10"
    wb.save(path)


def rebuild_all_orders_workbook():
    """Create a consolidated workbook containing every order and its line items."""
    with db_connection() as db:
        rows = db.execute("SELECT order_id, created_at, username, customer_name, phone, total, status, order_details FROM orders ORDER BY CAST(order_id AS INTEGER)").fetchall()
    wb = Workbook()
    summary = wb.active
    summary.title = "All Orders"
    summary.append(["Order number", "Date", "Account", "Customer", "Phone", "Total (INR)", "Status", "Product count"])
    details = wb.create_sheet("Order Items")
    details.append(["Order number", "Date", "Account", "Customer", "Product code", "Product", "Quantity", "Unit price (INR)", "Amount (INR)", "Status"])
    for row in rows:
        order = dict(row)
        data = json.loads(order["order_details"])
        summary.append([order["order_id"], order["created_at"].replace("T", " "), order["username"], order["customer_name"], order["phone"] or "", order["total"], order["status"], len(data.get("items", []))])
        for item in data.get("items", []):
            details.append([order["order_id"], order["created_at"].replace("T", " "), order["username"], order["customer_name"], item["code"], item["name"], item["quantity"], item["unit_price"], item["line_total"], order["status"]])
    style_workbook(summary); style_workbook(details)
    path = ALL_ORDERS_DIR / "all_orders.xlsx"
    wb.save(path)
    return path


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("username"):
            if request.path.startswith("/api/"):
                return jsonify({"success": False, "message": "Please log in to continue."}), 401
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)
    return wrapped

with PRODUCTS_FILE.open("r", encoding="utf-8") as f:
    PRODUCTS = json.load(f)


def db_connection():
    connection = sqlite3.connect(DB_FILE, timeout=30)
    connection.row_factory = sqlite3.Row
    return connection


def initialize_database():
    with db_connection() as db:
        db.execute("""
            CREATE TABLE IF NOT EXISTS orders (
                order_id TEXT PRIMARY KEY,
                created_at TEXT NOT NULL,
                customer_name TEXT NOT NULL,
                phone TEXT,
                total INTEGER NOT NULL,
                status TEXT NOT NULL,
                pdf_filename TEXT NOT NULL,
                order_details TEXT NOT NULL,
                username TEXT NOT NULL DEFAULT 'unknown'
            )
        """)
        # Migrate existing databases without losing previous orders.
        columns = {row[1] for row in db.execute("PRAGMA table_info(orders)").fetchall()}
        if "username" not in columns:
            db.execute("ALTER TABLE orders ADD COLUMN username TEXT NOT NULL DEFAULT 'unknown'")
        if "phone" not in columns:
            db.execute("ALTER TABLE orders ADD COLUMN phone TEXT")
        # Persistent counter ensures order references are 1, 2, 3... across restarts.
        db.execute("""
            CREATE TABLE IF NOT EXISTS order_sequence (
                sequence_name TEXT PRIMARY KEY,
                last_number INTEGER NOT NULL
            )
        """)
        db.execute("""
            INSERT OR IGNORE INTO order_sequence (sequence_name, last_number)
            VALUES ('orders', 0)
        """)
        db.execute("""
            CREATE TABLE IF NOT EXISTS inventory (
                product_id TEXT PRIMARY KEY,
                stock INTEGER NOT NULL DEFAULT 100,
                updated_at TEXT NOT NULL
            )
        """)
        # New products start at 100 units. Existing inventory is never overwritten on restart.
        now = datetime.now().astimezone().isoformat(timespec="seconds")
        for product in PRODUCTS:
            db.execute("INSERT OR IGNORE INTO inventory (product_id, stock, updated_at) VALUES (?, 100, ?)",
                       (str(product["id"]), now))


def reserve_order_number():
    # BEGIN IMMEDIATE serializes simultaneous order submissions to avoid duplicates.
    with db_connection() as db:
        db.execute("BEGIN IMMEDIATE")
        row = db.execute(
            "SELECT last_number FROM order_sequence WHERE sequence_name = 'orders'"
        ).fetchone()
        last_number = int(row["last_number"]) if row else 0
        number = last_number + 1
        db.execute(
            "INSERT INTO order_sequence (sequence_name, last_number) VALUES ('orders', ?) "
            "ON CONFLICT(sequence_name) DO UPDATE SET last_number = excluded.last_number",
            (number,)
        )
        db.commit()
        return number

initialize_database()


def escape_paragraph(value):
    return str(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_order_pdf(order):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ShopTitle", parent=styles["Title"], alignment=TA_CENTER,
                              textColor=colors.HexColor("#172554"), spaceAfter=8))
    styles.add(ParagraphStyle(name="Meta", parent=styles["Normal"], leading=17))
    story = [
        Paragraph("KKP FIREWORKS — ORDER RECEIPT", styles["ShopTitle"]),
        Paragraph("Order reference: " + escape_paragraph(order["order_id"]), styles["Meta"]),
        Paragraph("Date: " + escape_paragraph(order["created_at"].replace("T", " ")), styles["Meta"]),
        Paragraph("Mandal Name: " + escape_paragraph(order.get("mandal_name", order["customer_name"])), styles["Meta"]),
        Paragraph("Firm Name: " + escape_paragraph(order.get("firm_name", "Not provided")), styles["Meta"]),
        Paragraph("Mobile No: " + escape_paragraph(order.get("mobile_no", order.get("phone") or "Not provided")), styles["Meta"]),
        Spacer(1, 16),
    ]
    rows = [["Code", "Product", "Qty", "Unit price", "Amount"]]
    for item in order["items"]:
        rows.append([
            str(item["code"]), Paragraph(escape_paragraph(item["name"]), styles["BodyText"]),
            str(item["quantity"]), f'Rs. {item["unit_price"]:,}', f'Rs. {item["line_total"]:,}'
        ])
    rows.append(["", "", "", "TOTAL", f'Rs. {order["total"]:,}'])
    table = Table(rows, colWidths=[42, 220, 42, 75, 75], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#172554")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d8deea")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#f4f7fb")]),
        ("FONTNAME", (3, -1), (-1, -1), "Helvetica-Bold"),
        ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.extend([
        table, Spacer(1, 16),
        Paragraph("Order request only — confirm stock and final price with the customer.", styles["Italic"]),
        Spacer(1, 8), Paragraph("Created by Aditya Maganbhai Vasani", styles["Normal"])
    ])
    doc.build(story)
    return buffer.getvalue()


@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("username"):
        return redirect(url_for("home"))
    error = None
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        user = LOGIN_USERS.get(username)
        if user and secrets.compare_digest(password, user["password"]):
            session.clear()
            session["username"] = username
            session["role"] = user["role"]
            destination = request.args.get("next", "")
            # Only allow local relative redirects.
            if destination.startswith("/") and not destination.startswith("//"):
                return redirect(destination)
            return redirect(url_for("home"))
        error = "Invalid credentials. Please check your username and password."
    return render_template("login.html", error=error)

@app.get("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))

@app.get("/")
@login_required
def home():
    categories = list(dict.fromkeys(p["category"] for p in PRODUCTS))
    with db_connection() as db:
        stock_rows = db.execute("SELECT product_id, stock FROM inventory").fetchall()
    stock = {row["product_id"]: row["stock"] for row in stock_rows}
    products_with_stock = [dict(p, stock=stock.get(str(p["id"]), 100)) for p in PRODUCTS]
    return render_template("index.html", products=products_with_stock, categories=categories,
                           is_admin=session.get("role") == "admin")


@app.get("/api/products")
@login_required
def products():
    with db_connection() as db:
        stock_rows = db.execute("SELECT product_id, stock FROM inventory").fetchall()
    stock = {row["product_id"]: row["stock"] for row in stock_rows}
    items = [dict(p, stock=stock.get(str(p["id"]), 100)) for p in PRODUCTS]
    return jsonify({"success": True, "count": len(items), "products": items})


@app.post("/api/orders")
@login_required
def create_order():
    data = request.get_json(silent=True) or {}
    # Accept both the current Mandal/Firm/Mobile form fields and legacy clients.
    mandal_name = str(data.get("mandal_name", data.get("customer_name", ""))).strip()
    firm_name = str(data.get("firm_name", "")).strip()
    phone = str(data.get("mobile_no", data.get("phone", ""))).strip()
    customer_name = mandal_name
    items = data.get("items", [])

    if not mandal_name:
        return jsonify({"success": False, "message": "Please enter Mandal Name."}), 400
    if not firm_name:
        return jsonify({"success": False, "message": "Please enter Firm Name."}), 400
    if not re.fullmatch(r"[0-9]{10}", phone):
        return jsonify({"success": False, "message": "Please enter a valid 10-digit mobile number."}), 400
    if not isinstance(items, list) or not items:
        return jsonify({"success": False, "message": "Please select at least one product."}), 400

    product_by_id = {str(p["id"]): p for p in PRODUCTS}
    order_items, total = [], 0
    for item in items:
        if not isinstance(item, dict):
            continue
        product = product_by_id.get(str(item.get("product_id", item.get("id", ""))))
        try:
            quantity = int(item.get("quantity", 0))
        except (TypeError, ValueError):
            quantity = 0
        if not product or quantity == 0 or abs(quantity) > 1000000:
            continue
        line_total = int(product["price"]) * quantity
        total += line_total
        order_items.append({
            "code": product["code"], "product_id": product["id"], "name": product["name"],
            "category": product["category"], "unit": product["unit"], "quantity": quantity,
            "unit_price": int(product["price"]), "line_total": line_total
        })

    if not order_items:
        return jsonify({"success": False, "message": "No valid products were selected."}), 400

    now = datetime.now().astimezone()
    order_number = reserve_order_number()
    username = str(session.get("username", "unknown"))
    order = {
        "order_id": str(order_number),
        "username": username,
        "created_at": now.isoformat(timespec="seconds"),
        "customer_name": customer_name,
        "mandal_name": mandal_name,
        "firm_name": firm_name,
        "phone": phone,
        "mobile_no": phone,
        "items": order_items,
        "total": total,
        "status": "Order request received — confirm stock and final price with the customer."
    }
    pdf_bytes = build_order_pdf(order)
    pdf_filename = f'{order["order_id"]}_order.pdf'
    # Keep each user's order PDFs in their own directory on the persistent volume.
    user_folder_name = re.sub(r"[^a-zA-Z0-9_-]", "_", username) or "unknown"
    user_orders_dir = ORDERS_DIR / user_folder_name
    user_orders_dir.mkdir(parents=True, exist_ok=True)
    pdf_path = user_orders_dir / pdf_filename

    try:
        # Store both PDF and Excel receipt in the user's folder and a central all-orders folder.
        pdf_path.write_bytes(pdf_bytes)
        excel_filename = f'{order["order_id"]}_order.xlsx'
        excel_path = user_orders_dir / excel_filename
        build_order_excel(order, excel_path)
        all_pdf_path = ALL_ORDERS_DIR / pdf_filename
        all_excel_path = ALL_ORDERS_DIR / excel_filename
        shutil.copy2(pdf_path, all_pdf_path)
        shutil.copy2(excel_path, all_excel_path)
        # Store order history in SQLite on the same persistent Volume (no order.json files).
        with db_connection() as db:
            db.execute("""
                INSERT INTO orders (order_id, created_at, customer_name, phone, total, status, pdf_filename, order_details, username)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                order["order_id"], order["created_at"], customer_name, phone, total,
                order["status"], f"{user_folder_name}/{pdf_filename}", json.dumps(order, ensure_ascii=False), username
            ))
        rebuild_all_orders_workbook()
    except Exception:
        # Do not report success if durable storage failed.
        for cleanup_path in (pdf_path, user_orders_dir / f'{order_number}_order.xlsx', ALL_ORDERS_DIR / f'{order_number}_order.pdf', ALL_ORDERS_DIR / f'{order_number}_order.xlsx'):
            try:
                if cleanup_path.exists(): cleanup_path.unlink()
            except OSError:
                pass
        app.logger.exception("Failed to persist order")
        return jsonify({"success": False, "message": "Could not save the order. Please try again or contact the stall."}), 500

    import base64
    return jsonify({
        "success": True,
        "message": "Order saved successfully. Your PDF receipt is ready to download.",
        "order": order,
        "pdf_filename": pdf_filename,
        "excel_filename": excel_filename,
        "excel_url": url_for("download_my_order_excel", order_id=order["order_id"]),
        "pdf_base64": base64.b64encode(pdf_bytes).decode("ascii"),
        "saved_pdf_path": f"orders/{user_folder_name}/{pdf_filename}"
    }), 201


@app.get("/my-orders")
@login_required
def my_orders():
    # Admin can see every account's orders; mandal users see only their own.
    username = str(session.get("username", ""))
    is_admin = session.get("role") == "admin" or username == "admin"
    with db_connection() as db:
        if is_admin:
            rows = db.execute(
                "SELECT order_id, created_at, customer_name, phone, total, status, pdf_filename, order_details, username "
                "FROM orders ORDER BY CAST(order_id AS INTEGER) DESC"
            ).fetchall()
        else:
            rows = db.execute(
                "SELECT order_id, created_at, customer_name, phone, total, status, pdf_filename, order_details, username "
                "FROM orders WHERE username=? ORDER BY CAST(order_id AS INTEGER) DESC", (username,)
            ).fetchall()
    orders = []
    for row in rows:
        item = dict(row)
        try:
            parsed = json.loads(item.get("order_details") or "{}")
            if not isinstance(parsed, dict): parsed = {}
        except (TypeError, json.JSONDecodeError):
            parsed = {}
        parsed.setdefault("items", [])
        item["order_details"] = parsed
        orders.append(item)
    return render_template("my_orders.html", orders=orders, username=username, is_admin=is_admin)


@app.get("/my-orders/<order_id>/pdf")
@login_required
def download_my_order_pdf(order_id):
    username = str(session.get("username", ""))
    is_admin = session.get("role") == "admin" or username == "admin"
    with db_connection() as db:
        if is_admin:
            row = db.execute("SELECT pdf_filename FROM orders WHERE order_id=?", (str(order_id),)).fetchone()
        else:
            row = db.execute("SELECT pdf_filename FROM orders WHERE order_id=? AND username=?", (str(order_id), username)).fetchone()
    if not row:
        abort(404)
    path = ORDERS_DIR / row["pdf_filename"]
    if not path.is_file():
        abort(404)
    return send_file(path, as_attachment=True, download_name=f"{order_id}_order.pdf", mimetype="application/pdf")


@app.get("/my-orders/<order_id>/excel")
@login_required
def download_my_order_excel(order_id):
    username = str(session.get("username", ""))
    is_admin = session.get("role") == "admin" or username == "admin"
    with db_connection() as db:
        if is_admin:
            row = db.execute("SELECT order_details, username FROM orders WHERE order_id=?", (str(order_id),)).fetchone()
        else:
            row = db.execute("SELECT order_details, username FROM orders WHERE order_id=? AND username=?", (str(order_id), username)).fetchone()
    if not row:
        abort(404)
    order = json.loads(row["order_details"])
    order_owner = row["username"] if "username" in row.keys() else username
    folder = re.sub(r"[^a-zA-Z0-9_-]", "_", order_owner or "unknown") or "unknown"
    path = ORDERS_DIR / folder / f"{order_id}_order.xlsx"
    if not path.is_file():
        # Backfill Excel receipts for orders created before this feature.
        build_order_excel(order, path)
    return send_file(path, as_attachment=True, download_name=f"{order_id}_order.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/admin/all-orders.xlsx")
@admin_required
def download_all_orders_excel():
    path = rebuild_all_orders_workbook()
    return send_file(path, as_attachment=True, download_name="all_orders.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/admin/user-order-stock.xlsx")
@admin_required
def download_user_order_stock_excel():
    """Admin-only product matrix: quantities ordered by each account plus current stock."""
    customer_users = [username for username, info in LOGIN_USERS.items() if info.get("role") == "customer"]
    with db_connection() as db:
        stock_rows = db.execute("SELECT product_id, stock FROM inventory").fetchall()
        order_rows = db.execute("SELECT username, order_details FROM orders").fetchall()

    stock_by_id = {str(row["product_id"]): int(row["stock"]) for row in stock_rows}
    quantities = {username: {} for username in customer_users}
    for row in order_rows:
        username = row["username"] or "unknown"
        if username not in quantities:
            quantities[username] = {}
            customer_users.append(username)
        try:
            details = json.loads(row["order_details"] or "{}")
        except (TypeError, json.JSONDecodeError):
            continue
        for item in details.get("items", []):
            code = str(item.get("code", ""))
            try:
                qty = int(item.get("quantity", 0))
            except (TypeError, ValueError):
                qty = 0
            quantities[username][code] = quantities[username].get(code, 0) + qty

    wb = Workbook()
    ws = wb.active
    ws.title = "User Orders and Stock"
    headers = ["No.", "Item", "Product Code"] + customer_users + ["Total Ordered", "Remaining Stock"]
    ws.append(headers)
    for number, product in enumerate(PRODUCTS, start=1):
        code = str(product.get("code", ""))
        ordered_by_user = [quantities.get(username, {}).get(code, 0) for username in customer_users]
        total_ordered = sum(ordered_by_user)
        remaining = stock_by_id.get(str(product.get("id")), 100)
        ws.append([number, product.get("name", ""), code] + ordered_by_user + [total_ordered, remaining])
    style_workbook(ws)
    ws.freeze_panes = "D2"
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 16
    for col in range(4, len(headers) + 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 24 if col <= 3 + len(customer_users) else 18
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="kkp_user_orders_remaining_stock.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/admin/stock.xlsx")
@admin_required
def download_stock_excel():
    with db_connection() as db:
        rows = db.execute("SELECT product_id, stock FROM inventory").fetchall()
    stock_by_id = {row["product_id"]: int(row["stock"]) for row in rows}
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock Report"
    ws.append(["Product Code", "Product Name", "Category", "Current Stock", "Units To Replenish"])
    for product in PRODUCTS:
        current = stock_by_id.get(str(product["id"]), 100)
        ws.append([product["code"], product["name"], product["category"], current, max(0, -current)])
    style_workbook(ws)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="kkp_stock_report.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def build_whatsapp_url(phone, message):
    """Build a WhatsApp click-to-chat link from a customer mobile number."""
    digits = re.sub(r"\D", "", str(phone or ""))
    if len(digits) == 10:
        digits = "91" + digits
    if not digits:
        return ""
    return "https://wa.me/" + digits + "?text=" + quote_plus(message)


def attach_delivery_links(order):
    details = order.get("order_details", {})
    items = details.get("items", []) if isinstance(details, dict) else []
    lines = [
        f"Hello {order.get('customer_name') or 'Customer'},",
        f"This is regarding your KKP Fireworks order #{order.get('order_id')}.",
        "We are contacting you about delivery/collection of your order.",
        "Please reply to confirm a convenient time.",
        "Thank you, KKP Fireworks."
    ]
    order["whatsapp_url"] = build_whatsapp_url(order.get("phone"), "\n".join(lines))
    delivery_lines = [
        f"Hello {order.get('customer_name') or 'Customer'},",
        f"Your KKP Fireworks order #{order.get('order_id')} is being arranged for delivery/collection.",
        "Items: " + (", ".join(f"{i.get('name', i.get('code', 'Item'))} x {i.get('quantity', 0)}" for i in items) or "As per your order"),
        f"Order total: Rs. {int(order.get('total') or 0):,}",
        "Please reply to coordinate delivery. Thank you, KKP Fireworks."
    ]
    order["delivery_whatsapp_url"] = build_whatsapp_url(order.get("phone"), "\n".join(delivery_lines))
    return order


@app.get("/admin")
@admin_required
def admin_dashboard():
    with db_connection() as db:
        rows = db.execute("SELECT order_id, created_at, customer_name, phone, total, status, order_details, username FROM orders ORDER BY CAST(order_id AS INTEGER) DESC").fetchall()
        stock_rows = db.execute("SELECT product_id, stock FROM inventory").fetchall()
    stock = {row["product_id"]: row["stock"] for row in stock_rows}
    inventory = [dict(p, stock=stock.get(str(p["id"]), 100)) for p in PRODUCTS]
    orders = []
    for row in rows:
        item = dict(row)
        # A malformed historical order must not crash the entire admin dashboard.
        try:
            parsed = json.loads(item.get("order_details") or "{}")
            if not isinstance(parsed, dict):
                parsed = {}
        except (TypeError, json.JSONDecodeError):
            parsed = {}
        parsed.setdefault("items", [])
        item["order_details"] = parsed
        item.setdefault("username", "unknown")
        attach_delivery_links(item)
        orders.append(item)
    # Show all known customer accounts in the admin-only per-user report selector.
    report_users = sorted({
        username for username, info in LOGIN_USERS.items() if info.get("role") == "customer"
    } | {str(order.get("username") or "unknown") for order in orders if order.get("username")})
    return render_template("admin.html", inventory=inventory, orders=orders, report_users=report_users)


def _get_user_orders_for_report(username):
    """Return one account's orders, parsed and ordered for admin report exports."""
    allowed_users = {u for u, info in LOGIN_USERS.items() if info.get("role") == "customer"}
    with db_connection() as db:
        known_rows = db.execute("SELECT DISTINCT username FROM orders WHERE username IS NOT NULL AND username != ''").fetchall()
        allowed_users.update(str(row["username"]) for row in known_rows)
        if username not in allowed_users:
            abort(404, description="Customer account not found.")
        rows = db.execute(
            "SELECT order_id, created_at, username, customer_name, phone, total, status, order_details "
            "FROM orders WHERE username=? ORDER BY CAST(order_id AS INTEGER), order_id", (username,)
        ).fetchall()
    orders_for_user = []
    for row in rows:
        record = dict(row)
        try:
            details = json.loads(record.get("order_details") or "{}")
            if not isinstance(details, dict):
                details = {}
        except (TypeError, json.JSONDecodeError):
            details = {}
        record["details"] = details
        record["items"] = details.get("items", []) if isinstance(details.get("items", []), list) else []
        record["mandal_name"] = details.get("mandal_name") or record.get("customer_name") or ""
        record["firm_name"] = details.get("firm_name") or ""
        record["phone"] = details.get("mobile_no") or record.get("phone") or ""
        try:
            record["total"] = int(record.get("total") or details.get("total") or 0)
        except (TypeError, ValueError):
            record["total"] = 0
        orders_for_user.append(record)
    return orders_for_user


@app.get("/admin/user-orders.xlsx")
@admin_required
def download_one_user_orders_excel():
    username = str(request.args.get("username", "")).strip()
    if not username:
        abort(400, description="Select a customer account first.")
    orders_for_user = _get_user_orders_for_report(username)
    wb = Workbook()
    ws = wb.active
    ws.title = "User Order Report"
    ws.append(["KKP FIREWORKS — CUSTOMER ORDER REPORT"])
    ws.merge_cells("A1:K1")
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="172554")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(["Login account", username])
    ws.append(["Orders included", len(orders_for_user)])
    ws.append([])
    ws.append(["Order No.", "Date", "Status", "Mandal / Customer", "Firm Name", "Mobile", "Product Code", "Product", "Quantity", "Unit Price (INR)", "Line Total (INR)"])
    for cell in ws[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="172554")
    grand_total = 0
    for order in orders_for_user:
        details = order["details"]
        for item in order["items"]:
            try: qty = int(item.get("quantity", 0))
            except (TypeError, ValueError): qty = 0
            try: price = int(item.get("unit_price", 0))
            except (TypeError, ValueError): price = 0
            try: amount = int(item.get("line_total", qty * price))
            except (TypeError, ValueError): amount = qty * price
            grand_total += amount
            ws.append([order["order_id"], order.get("created_at", "").replace("T", " "), order.get("status", ""), order["mandal_name"], order["firm_name"], order["phone"], item.get("code", ""), item.get("name", ""), qty, price, amount])
    ws.append([])
    ws.append(["", "", "", "", "", "", "", "", "", "Grand total", grand_total])
    style_workbook(ws)
    # Restore the report title/header styling after general workbook styling.
    for cell in ws[1]:
        cell.font = Font(bold=True, size=14, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="172554")
    for cell in ws[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="172554")
    ws.freeze_panes = "A6"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 24
    ws.column_dimensions["E"].width = 24
    ws.column_dimensions["F"].width = 15
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 38
    for col in ("I", "J", "K"): ws.column_dimensions[col].width = 18
    for row in ws.iter_rows(min_row=6):
        for cell in row:
            if cell.column in (10, 11) and isinstance(cell.value, (int, float)):
                cell.number_format = '₹#,##0'
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    safe_name = re.sub(r"[^a-zA-Z0-9_-]", "_", username) or "customer"
    return send_file(output, as_attachment=True, download_name=f"{safe_name}_orders_report.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/admin/user-orders.pdf")
@admin_required
def download_one_user_orders_pdf():
    username = str(request.args.get("username", "")).strip()
    if not username:
        abort(400, description="Select a customer account first.")
    orders_for_user = _get_user_orders_for_report(username)
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=28, leftMargin=28, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="UserReportTitle", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#172554"), spaceAfter=10))
    story = [Paragraph("KKP FIREWORKS — CUSTOMER ORDER REPORT", styles["UserReportTitle"]),
             Paragraph("Login account: " + escape_paragraph(username), styles["Normal"]),
             Paragraph("Orders: " + str(len(orders_for_user)), styles["Normal"]), Spacer(1, 10)]
    table_rows = [["Order", "Date", "Status", "Customer / Mandal", "Firm", "Mobile", "Product Code", "Product", "Qty", "Unit Price", "Amount"]]
    grand_total = 0
    for order in orders_for_user:
        for item in order["items"]:
            try: qty = int(item.get("quantity", 0))
            except (TypeError, ValueError): qty = 0
            try: price = int(item.get("unit_price", 0))
            except (TypeError, ValueError): price = 0
            try: amount = int(item.get("line_total", qty * price))
            except (TypeError, ValueError): amount = qty * price
            grand_total += amount
            table_rows.append([str(order["order_id"]), str(order.get("created_at", "")).replace("T", " ")[:16], str(order.get("status", ""))[:18], str(order["mandal_name"])[:28], str(order["firm_name"])[:24], str(order["phone"]), str(item.get("code", "")), str(item.get("name", ""))[:36], str(qty), f"Rs. {price:,}", f"Rs. {amount:,}"])
    table_rows.append(["", "", "", "", "", "", "", "GRAND TOTAL", "", "", f"Rs. {grand_total:,}"])
    table = Table(table_rows, repeatRows=1, colWidths=[38, 68, 58, 85, 70, 55, 48, 115, 30, 55, 65])
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#172554")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("FONTSIZE", (0,0), (-1,-1), 7), ("GRID", (0,0), (-1,-1), 0.35, colors.HexColor("#d8deea")), ("ROWBACKGROUNDS", (0,1), (-1,-2), [colors.white, colors.HexColor("#f4f7fb")]), ("VALIGN", (0,0), (-1,-1), "TOP"), ("ALIGN", (8,1), (-1,-1), "RIGHT")]))
    story.extend([table, Spacer(1, 12), Paragraph(f"Grand total across line items: Rs. {grand_total:,}", styles["Heading3"]), Spacer(1, 8), Paragraph("Created by Aditya Maganbhai Vasani", styles["Normal"])])
    doc.build(story)
    buffer.seek(0)
    safe_name = re.sub(r"[^a-zA-Z0-9_-]", "_", username) or "customer"
    return send_file(buffer, as_attachment=True, download_name=f"{safe_name}_orders_report.pdf", mimetype="application/pdf")


@app.get("/admin/delivery-report")
@admin_required
def admin_delivery_report():
    """Admin-only delivery report; reads existing orders without modifying them."""
    with db_connection() as db:
        rows = db.execute(
            "SELECT order_id, created_at, customer_name, phone, total, status, order_details, username "
            "FROM orders ORDER BY CAST(order_id AS INTEGER) DESC"
        ).fetchall()
    orders = []
    for row in rows:
        item = dict(row)
        try:
            parsed = json.loads(item.get("order_details") or "{}")
            if not isinstance(parsed, dict): parsed = {}
        except (TypeError, json.JSONDecodeError):
            parsed = {}
        parsed.setdefault("items", [])
        item["order_details"] = parsed
        item.setdefault("username", "unknown")
        attach_delivery_links(item)
        orders.append(item)
    return render_template("delivery_report.html", orders=orders)


@app.post("/api/admin/stock")
@admin_required
def update_stock():
    data = request.get_json(silent=True) or {}
    product_id = str(data.get("product_id", ""))
    try:
        stock = int(data.get("stock"))
    except (TypeError, ValueError):
        return jsonify({"success": False, "message": "Stock must be a whole number."}), 400
    if stock < -1000000 or stock > 1000000:
        return jsonify({"success": False, "message": "Stock must be between -1,000,000 and 1,000,000."}), 400
    if product_id not in {str(p["id"]) for p in PRODUCTS}:
        return jsonify({"success": False, "message": "Unknown product."}), 404
    with db_connection() as db:
        db.execute("UPDATE inventory SET stock=?, updated_at=? WHERE product_id=?",
                   (stock, datetime.now().astimezone().isoformat(timespec="seconds"), product_id))
    return jsonify({"success": True, "message": "Stock updated.", "stock": stock})


@app.post("/api/admin/orders/<order_id>/change-number")
@admin_required
def change_order_number(order_id):
    """Change an order's public reference without deleting its history or changing stock."""
    data = request.get_json(silent=True) or {}
    new_order_id = str(data.get("new_order_id", "")).strip()
    if not re.fullmatch(r"[0-9]{1,12}", new_order_id) or int(new_order_id) < 1:
        return jsonify({"success": False, "message": "Enter a positive order number using digits only (1–12 digits)."}), 400
    new_order_id = str(int(new_order_id))  # Canonicalize so 01 and 1 cannot be used as separate references.
    old_order_id = str(order_id)
    if new_order_id == old_order_id:
        return jsonify({"success": False, "message": "The new order number is the same as the current number."}), 400

    with db_connection() as db:
        row = db.execute(
            "SELECT order_id, created_at, customer_name, phone, total, status, pdf_filename, order_details, username "
            "FROM orders WHERE order_id=?", (old_order_id,)
        ).fetchone()
        if not row:
            return jsonify({"success": False, "message": "Order not found."}), 404
        if db.execute("SELECT 1 FROM orders WHERE order_id=?", (new_order_id,)).fetchone():
            return jsonify({"success": False, "message": "That order number is already in use. Choose a unique number."}), 409

        order = dict(row)
        try:
            details = json.loads(order.get("order_details") or "{}")
            if not isinstance(details, dict):
                details = {}
        except (TypeError, json.JSONDecodeError):
            details = {}
        details.setdefault("items", [])
        details["order_id"] = new_order_id
        details["status"] = order["status"]
        details.setdefault("customer_name", order["customer_name"])
        details.setdefault("phone", order["phone"] or "")
        details.setdefault("total", order["total"])
        details.setdefault("created_at", order["created_at"])
        details.setdefault("username", order["username"] or "unknown")

        username_folder = re.sub(r"[^a-zA-Z0-9_-]", "_", order["username"] or "unknown") or "unknown"
        user_dir = ORDERS_DIR / username_folder
        user_dir.mkdir(parents=True, exist_ok=True)
        old_pdf_name = order["pdf_filename"]
        new_pdf_name = f"{new_order_id}_order.pdf"
        new_excel_name = f"{new_order_id}_order.xlsx"
        new_pdf_path = user_dir / new_pdf_name
        new_excel_path = user_dir / new_excel_name
        all_new_pdf = ALL_ORDERS_DIR / new_pdf_name
        all_new_excel = ALL_ORDERS_DIR / new_excel_name
        old_user_pdf = ORDERS_DIR / old_pdf_name
        old_user_excel = user_dir / f"{old_order_id}_order.xlsx"
        all_old_pdf = ALL_ORDERS_DIR / f"{old_order_id}_order.pdf"
        all_old_excel = ALL_ORDERS_DIR / f"{old_order_id}_order.xlsx"
        if any(path.exists() for path in (new_pdf_path, new_excel_path, all_new_pdf, all_new_excel)):
            return jsonify({"success": False, "message": "Files already exist for that order number. Choose another number."}), 409

        # Build replacement receipts first so they display the new number too.
        temp_token = secrets.token_hex(6)
        temp_pdf = user_dir / f".{new_order_id}_{temp_token}.pdf.tmp"
        temp_excel = user_dir / f".{new_order_id}_{temp_token}.xlsx.tmp"
        try:
            temp_pdf.write_bytes(build_order_pdf(details))
            build_order_excel(details, temp_excel)
            # Install the updated receipts before committing the reference change.
            os.replace(temp_pdf, new_pdf_path)
            os.replace(temp_excel, new_excel_path)
            shutil.copy2(new_pdf_path, all_new_pdf)
            shutil.copy2(new_excel_path, all_new_excel)
            with db:
                db.execute("UPDATE orders SET order_id=?, pdf_filename=?, order_details=? WHERE order_id=?",
                           (new_order_id, f"{username_folder}/{new_pdf_name}", json.dumps(details, ensure_ascii=False), old_order_id))
                # Keep the automatic counter above the largest manually assigned number.
                db.execute("UPDATE order_sequence SET last_number=MAX(last_number, ?) WHERE sequence_name='orders'",
                           (int(new_order_id),))
            # Remove superseded receipt copies only after the database update succeeds.
            for path in (old_user_pdf, old_user_excel, all_old_pdf, all_old_excel):
                try:
                    if path.exists() and path not in (new_pdf_path, new_excel_path):
                        path.unlink()
                except OSError:
                    app.logger.warning("Could not remove old receipt file: %s", path)
            try:
                rebuild_all_orders_workbook()
            except Exception:
                app.logger.exception("Order number changed, but consolidated workbook refresh failed")
        except Exception:
            for path in (temp_pdf, temp_excel, new_pdf_path, new_excel_path, all_new_pdf, all_new_excel):
                try:
                    if path.exists(): path.unlink()
                except OSError:
                    pass
            app.logger.exception("Failed to change order number %s to %s", old_order_id, new_order_id)
            return jsonify({"success": False, "message": "Could not update the order number. Check server logs before retrying."}), 500

    return jsonify({"success": True, "message": f"Order number changed from #{old_order_id} to #{new_order_id}.", "order_id": new_order_id})


@app.post("/api/admin/orders/<order_id>/confirm")
@admin_required
def confirm_order(order_id):
    # Apply signed stock movement only when admin confirms: positive quantities reduce stock; negative quantities increase it.
    with db_connection() as db:
        db.execute("BEGIN IMMEDIATE")
        row = db.execute("SELECT status, order_details FROM orders WHERE order_id=?", (str(order_id),)).fetchone()
        if not row:
            return jsonify({"success": False, "message": "Order not found."}), 404
        if row["status"] == "Confirmed":
            return jsonify({"success": False, "message": "This order is already confirmed."}), 409
        order = json.loads(row["order_details"])
        # Signed quantity is applied as stock = stock - quantity. A negative quantity therefore returns stock.
        for item in order["items"]:
            db.execute("UPDATE inventory SET stock=stock-?, updated_at=? WHERE product_id=?",
                       (int(item["quantity"]), datetime.now().astimezone().isoformat(timespec="seconds"), str(item["product_id"])))
        order["status"] = "Confirmed"
        db.execute("UPDATE orders SET status=?, order_details=? WHERE order_id=?",
                   ("Confirmed", json.dumps(order, ensure_ascii=False), str(order_id)))
        db.commit()
    return jsonify({"success": True, "message": "Order confirmed and stock updated using the signed quantities."})


@app.post("/api/admin/orders/<order_id>/edit-pending")
@admin_required
def edit_pending_order(order_id):
    """Edit item quantities only while an order is still awaiting admin confirmation."""
    data = request.get_json(silent=True) or {}
    quantities = data.get("quantities")
    if not isinstance(quantities, dict):
        return jsonify({"success": False, "message": "Provide a quantity for each order item."}), 400

    # Read and validate the current order before writing anything.
    with db_connection() as db:
        row = db.execute(
            "SELECT order_id, created_at, customer_name, phone, total, status, pdf_filename, order_details, username FROM orders WHERE order_id=?",
            (str(order_id),)
        ).fetchone()
        if not row:
            return jsonify({"success": False, "message": "Order not found. Refresh the admin page and try again."}), 404
        if str(row["status"] or "").lower().startswith("confirmed"):
            return jsonify({"success": False, "message": "Confirmed orders cannot be edited."}), 409
        try:
            details = json.loads(row["order_details"] or "{}")
            items = details.get("items", [])
            if not isinstance(items, list) or not items:
                raise ValueError("No items")
        except (TypeError, ValueError, json.JSONDecodeError):
            return jsonify({"success": False, "message": "Order items could not be read."}), 400

        updated_items = []
        total = 0
        for index, item in enumerate(items):
            raw_qty = quantities.get(str(index))
            try:
                qty = int(raw_qty)
            except (TypeError, ValueError):
                return jsonify({"success": False, "message": f"Invalid quantity for {item.get('name', 'item')}."}), 400
            if qty == 0 or abs(qty) > 1000000:
                return jsonify({"success": False, "message": "Each quantity must be a non-zero whole number between -1,000,000 and 1,000,000."}), 400
            try:
                unit_price = int(item.get("unit_price", 0))
            except (TypeError, ValueError):
                return jsonify({"success": False, "message": f"Invalid saved price for {item.get('name', 'item')}."}), 400
            item = dict(item)
            item["quantity"] = qty
            item["line_total"] = qty * unit_price
            total += item["line_total"]
            updated_items.append(item)

        details["items"] = updated_items
        details["total"] = total
        details["status"] = row["status"]
        details["order_id"] = str(order_id)
        details["created_at"] = details.get("created_at") or row["created_at"]
        details["customer_name"] = details.get("customer_name") or row["customer_name"]
        details["phone"] = details.get("phone") or row["phone"] or ""
        details["username"] = details.get("username") or row["username"] or "unknown"

        # Save the order in SQLite first. Receipt generation is refreshed afterward so a
        # missing/corrupt receipt file cannot prevent the actual order edit from saving.
        try:
            db.execute("UPDATE orders SET total=?, order_details=? WHERE order_id=?",
                       (total, json.dumps(details, ensure_ascii=False), str(order_id)))
            db.commit()
        except Exception:
            db.rollback()
            app.logger.exception("Could not save pending order edit %s", order_id)
            return jsonify({"success": False, "message": "Could not save the edited order to the database."}), 500

    receipt_warning = ""
    try:
        username_folder = re.sub(r"[^a-zA-Z0-9_-]", "_", row["username"] or "unknown") or "unknown"
        user_dir = ORDERS_DIR / username_folder
        user_dir.mkdir(parents=True, exist_ok=True)
        pdf_name = f"{order_id}_order.pdf"
        excel_name = f"{order_id}_order.xlsx"
        user_pdf = user_dir / pdf_name
        user_excel = user_dir / excel_name
        temp_pdf = user_dir / f".{order_id}_edit.tmp.pdf"
        temp_excel = user_dir / f".{order_id}_edit.tmp.xlsx"
        temp_pdf.write_bytes(build_order_pdf(details))
        build_order_excel(details, temp_excel)
        os.replace(temp_pdf, user_pdf)
        os.replace(temp_excel, user_excel)
        shutil.copy2(user_pdf, ALL_ORDERS_DIR / pdf_name)
        shutil.copy2(user_excel, ALL_ORDERS_DIR / excel_name)
        # Keep the existing stored relative PDF path, repairing it only if needed.
        stored_pdf = f"{username_folder}/{pdf_name}"
        with db_connection() as db:
            db.execute("UPDATE orders SET pdf_filename=? WHERE order_id=?", (stored_pdf, str(order_id)))
            db.commit()
    except Exception:
        receipt_warning = " The order was saved, but its receipt files could not be refreshed; check the server logs."
        app.logger.exception("Order %s was saved but receipt refresh failed", order_id)
    try:
        rebuild_all_orders_workbook()
    except Exception:
        app.logger.exception("Order %s was saved but consolidated workbook refresh failed", order_id)
    return jsonify({"success": True, "message": f"Order #{order_id} quantities saved. Updated total: ₹{total:,}.{receipt_warning}", "total": total})


@app.post("/api/admin/orders/<order_id>/edit-confirmed")
@admin_required
def edit_confirmed_order(order_id):
    """Edit a confirmed order and reconcile inventory by the quantity difference."""
    data = request.get_json(silent=True) or {}
    quantities = data.get("quantities")
    if not isinstance(quantities, dict):
        return jsonify({"success": False, "message": "Provide a quantity for each order item."}), 400

    now = datetime.now().astimezone().isoformat(timespec="seconds")
    with db_connection() as db:
        db.execute("BEGIN IMMEDIATE")
        row = db.execute(
            "SELECT order_id, created_at, customer_name, phone, total, status, pdf_filename, order_details, username FROM orders WHERE order_id=?",
            (str(order_id),)
        ).fetchone()
        if not row:
            db.rollback()
            return jsonify({"success": False, "message": "Order not found."}), 404
        if str(row["status"] or "").strip().lower() != "confirmed":
            db.rollback()
            return jsonify({"success": False, "message": "This order is not confirmed. Use Edit quantities for pending orders."}), 409
        try:
            details = json.loads(row["order_details"] or "{}")
            items = details.get("items", [])
            if not isinstance(items, list) or not items:
                raise ValueError("No items")
            updated_items = []
            total = 0
            stock_changes = []
            for index, old_item in enumerate(items):
                product_id = str(old_item.get("product_id", ""))
                old_qty = int(old_item.get("quantity", 0))
                if not product_id or old_qty == 0 or abs(old_qty) > 1000000:
                    raise ValueError("Invalid saved item quantity")
                raw_qty = quantities.get(str(index))
                new_qty = int(raw_qty)
                if new_qty == 0 or abs(new_qty) > 1000000:
                    raise ValueError("Each quantity must be a non-zero whole number between -1,000,000 and 1,000,000")
                unit_price = int(old_item.get("unit_price", 0))
                if unit_price < 0:
                    raise ValueError("Invalid saved price")
                item = dict(old_item)
                item["quantity"] = new_qty
                item["line_total"] = new_qty * unit_price
                total += item["line_total"]
                updated_items.append(item)
                # Stock was already reduced by old_qty at confirmation. Reconcile only the delta.
                stock_changes.append((product_id, old_qty - new_qty))
            details["items"] = updated_items
            details["total"] = total
            details["status"] = row["status"]
            details["order_id"] = str(order_id)
            details["created_at"] = details.get("created_at") or row["created_at"]
            details["customer_name"] = details.get("customer_name") or row["customer_name"]
            details["phone"] = details.get("phone") or row["phone"] or ""
            details["username"] = details.get("username") or row["username"] or "unknown"
            for product_id, delta in stock_changes:
                if delta:
                    cur = db.execute("UPDATE inventory SET stock=stock+?, updated_at=? WHERE product_id=?", (delta, now, product_id))
                    if cur.rowcount == 0:
                        raise ValueError(f"Inventory record missing for product {product_id}")
            db.execute("UPDATE orders SET total=?, order_details=? WHERE order_id=?", (total, json.dumps(details, ensure_ascii=False), str(order_id)))
            db.commit()
        except (TypeError, ValueError, json.JSONDecodeError) as exc:
            db.rollback()
            return jsonify({"success": False, "message": f"Order was not changed: {exc}."}), 400
        except Exception:
            db.rollback()
            app.logger.exception("Could not save confirmed order edit %s", order_id)
            return jsonify({"success": False, "message": "Could not save the edit. No stock or order changes were committed."}), 500

    receipt_warning = ""
    try:
        username_folder = re.sub(r"[^a-zA-Z0-9_-]", "_", row["username"] or "unknown") or "unknown"
        user_dir = ORDERS_DIR / username_folder
        user_dir.mkdir(parents=True, exist_ok=True)
        pdf_name = f"{order_id}_order.pdf"
        excel_name = f"{order_id}_order.xlsx"
        user_pdf, user_excel = user_dir / pdf_name, user_dir / excel_name
        temp_pdf, temp_excel = user_dir / f".{order_id}_edit.tmp.pdf", user_dir / f".{order_id}_edit.tmp.xlsx"
        temp_pdf.write_bytes(build_order_pdf(details))
        build_order_excel(details, temp_excel)
        os.replace(temp_pdf, user_pdf)
        os.replace(temp_excel, user_excel)
        shutil.copy2(user_pdf, ALL_ORDERS_DIR / pdf_name)
        shutil.copy2(user_excel, ALL_ORDERS_DIR / excel_name)
        stored_pdf = f"{username_folder}/{pdf_name}"
        with db_connection() as db:
            db.execute("UPDATE orders SET pdf_filename=? WHERE order_id=?", (stored_pdf, str(order_id)))
            db.commit()
    except Exception:
        receipt_warning = " The order and stock were updated, but receipt refresh failed; check server logs."
        app.logger.exception("Confirmed order %s saved but receipt refresh failed", order_id)
    try:
        rebuild_all_orders_workbook()
    except Exception:
        app.logger.exception("Confirmed order %s saved but consolidated workbook refresh failed", order_id)
    return jsonify({"success": True, "message": f"Confirmed order #{order_id} updated. Stock was reconciled by the quantity difference. New total: ₹{total:,}.{receipt_warning}", "total": total})


@app.post("/api/admin/orders/<order_id>/delete-pending")
@admin_required
def delete_pending_order(order_id):
    """Delete only unconfirmed/fake order requests; confirmed orders are protected."""
    with db_connection() as db:
        row = db.execute(
            "SELECT order_id, status, pdf_filename, username FROM orders WHERE order_id=?", (str(order_id),)
        ).fetchone()
        if not row:
            return jsonify({"success": False, "message": "Order not found."}), 404
        if str(row["status"]).lower() == "confirmed":
            return jsonify({"success": False, "message": "Confirmed orders cannot be deleted."}), 409
        db.execute("DELETE FROM orders WHERE order_id=? AND status != 'Confirmed'", (str(order_id),))
        if db.total_changes == 0:
            db.rollback()
            return jsonify({"success": False, "message": "Order was not deleted."}), 409
        db.commit()

    username_folder = re.sub(r"[^a-zA-Z0-9_-]", "_", row["username"] or "unknown") or "unknown"
    paths = [ORDERS_DIR / row["pdf_filename"], ORDERS_DIR / username_folder / f"{order_id}_order.xlsx",
             ALL_ORDERS_DIR / f"{order_id}_order.pdf", ALL_ORDERS_DIR / f"{order_id}_order.xlsx"]
    for path in paths:
        try:
            if path.is_file():
                path.unlink()
        except OSError:
            app.logger.warning("Could not remove receipt for deleted pending order: %s", path)
    try:
        rebuild_all_orders_workbook()
    except Exception:
        app.logger.exception("Deleted pending order but consolidated workbook refresh failed")
    return jsonify({"success": True, "message": f"Unconfirmed order #{order_id} deleted. Stock was not changed."})


@app.post("/api/admin/orders/<order_id>/delete-confirmed")
@admin_required
def delete_confirmed_order(order_id):
    """Admin-only recovery for an order confirmed by mistake. Restore deducted stock and remove order/receipts."""
    with db_connection() as db:
        db.execute("BEGIN IMMEDIATE")
        row = db.execute(
            "SELECT order_id, status, pdf_filename, username, order_details FROM orders WHERE order_id=?",
            (str(order_id),)
        ).fetchone()
        if not row:
            db.rollback()
            return jsonify({"success": False, "message": "Order not found."}), 404
        if str(row["status"] or "").strip().lower() != "confirmed":
            db.rollback()
            return jsonify({"success": False, "message": "This order is not confirmed. Use the pending-order delete action instead."}), 409
        try:
            details = json.loads(row["order_details"] or "{}")
            items = details.get("items", [])
            if not isinstance(items, list):
                raise ValueError("Invalid items")
            # Reverse the original signed stock movement in the same transaction (stock += quantity).
            now = datetime.now().astimezone().isoformat(timespec="seconds")
            for item in items:
                product_id = str(item.get("product_id", ""))
                quantity = int(item.get("quantity", 0))
                if not product_id or abs(quantity) > 1000000:
                    raise ValueError("Invalid product quantity")
                if quantity:
                    cur = db.execute(
                        "UPDATE inventory SET stock=stock+?, updated_at=? WHERE product_id=?",
                        (quantity, now, product_id)
                    )
                    if cur.rowcount == 0:
                        raise ValueError(f"Inventory record missing for product {product_id}")
        except (TypeError, ValueError, json.JSONDecodeError) as exc:
            db.rollback()
            app.logger.exception("Could not safely reverse stock for confirmed order %s", order_id)
            return jsonify({"success": False, "message": "Could not safely delete this order because its saved item quantities could not be validated. No changes were made."}), 409
        db.execute("DELETE FROM orders WHERE order_id=?", (str(order_id),))
        db.commit()

    username_folder = re.sub(r"[^a-zA-Z0-9_-]", "_", row["username"] or "unknown") or "unknown"
    stored_pdf = str(row["pdf_filename"] or "")
    paths = [ORDERS_DIR / stored_pdf, ORDERS_DIR / username_folder / f"{order_id}_order.pdf",
             ORDERS_DIR / username_folder / f"{order_id}_order.xlsx",
             ALL_ORDERS_DIR / f"{order_id}_order.pdf", ALL_ORDERS_DIR / f"{order_id}_order.xlsx"]
    for path in paths:
        try:
            if path.is_file():
                path.unlink()
        except OSError:
            app.logger.warning("Could not remove receipt for deleted confirmed order: %s", path)
    try:
        rebuild_all_orders_workbook()
    except Exception:
        app.logger.exception("Deleted confirmed order %s but consolidated workbook refresh failed", order_id)
    return jsonify({"success": True, "message": f"Confirmed order #{order_id} deleted. Its deducted quantities were added back to stock."})


@app.get("/health")
def health():
    return jsonify({"status": "ok", "products": len(PRODUCTS)})


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5000"))
    app.run(host="0.0.0.0", port=port, debug=False)
